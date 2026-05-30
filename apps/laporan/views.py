from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q, Sum
from datetime import datetime, timedelta
from calendar import month_name, monthrange
from datetime import date
from apps.balita.models import Balita
from apps.penimbangan.models import Penimbangan
from apps.imunisasi.models import Imunisasi
from django.http import HttpResponse
import json


@login_required
def laporan_index(request):
    """Halaman laporan bulanan dan tahunan"""
    # Get current month and year
    now = datetime.now()
    jenis_laporan = request.GET.get('jenis', 'bulanan')  # 'bulanan' or 'tahunan'
    bulan = int(request.GET.get('bulan', now.month))
    tahun = int(request.GET.get('tahun', now.year))
    
    if jenis_laporan == 'tahunan':
        return laporan_tahunan(request, tahun)
    
    # ==================== LAPORAN BULANAN ====================
    # Filter date range
    start_date = date(tahun, bulan, 1)
    last_day = monthrange(tahun, bulan)[1]
    end_date = date(tahun, bulan, last_day)
    
    # Total balita (seluruhnya)
    total_balita = Balita.objects.count()
    
    # Total balita baru bulan ini
    balita_baru = Balita.objects.filter(created_at__date__gte=start_date, created_at__date__lte=end_date).count()
    
    # Data penimbangan bulan ini
    penimbangan_list = Penimbangan.objects.filter(tanggal__gte=start_date, tanggal__lte=end_date)
    total_penimbangan = penimbangan_list.count()
    
    # Data imunisasi bulan ini
    imunisasi_list = Imunisasi.objects.filter(tanggal__gte=start_date, tanggal__lte=end_date)
    total_imunisasi = imunisasi_list.count()
    
    # Persentase balita yang ditimbang
    balita_ditimbang = penimbangan_list.values('balita').distinct().count()
    persentase_penimbangan = round((balita_ditimbang / total_balita * 100) if total_balita > 0 else 0)
    
    # Status gizi bulan ini
    status_normal = 0
    status_kurang = 0
    status_buruk = 0
    status_lebih = 0
    
    for balita in Balita.objects.all():
        penimbangan_terakhir = balita.penimbangan.filter(tanggal__lte=end_date).order_by('-tanggal').first()
        if penimbangan_terakhir and penimbangan_terakhir.status_gizi:
            status = penimbangan_terakhir.status_gizi
            if 'Normal' in status:
                status_normal += 1
            elif 'Kurang' in status:
                status_kurang += 1
            elif 'Buruk' in status:
                status_buruk += 1
            elif 'Lebih' in status:
                status_lebih += 1
    
    total_status = status_normal + status_kurang + status_buruk + status_lebih
    
    persen_normal = round((status_normal / total_status * 100), 1) if total_status > 0 else 0
    persen_kurang = round((status_kurang / total_status * 100), 1) if total_status > 0 else 0
    persen_buruk = round((status_buruk / total_status * 100), 1) if total_status > 0 else 0
    persen_lebih = round((status_lebih / total_status * 100), 1) if total_status > 0 else 0
    
    # Chart data untuk grafik penimbangan per hari
    chart_labels = []
    chart_data = []
    for day in range(1, last_day + 1):
        tgl = date(tahun, bulan, day)
        chart_labels.append(str(day))
        count = Penimbangan.objects.filter(tanggal=tgl).count()
        chart_data.append(count)
    
    # Data untuk perbandingan antar bulan (6 bulan terakhir)
    bulan_labels = []
    bulan_penimbangan = []
    bulan_imunisasi = []
    for i in range(5, -1, -1):
        bulan_tgl = date(tahun, bulan, 1) - timedelta(days=30 * i)
        bln = bulan_tgl.month
        thn = bulan_tgl.year
        start = date(thn, bln, 1)
        last = date(thn, bln, monthrange(thn, bln)[1])
        bulan_labels.append(f"{month_name[bln][:3]} {thn}")
        bulan_penimbangan.append(Penimbangan.objects.filter(tanggal__gte=start, tanggal__lte=last).count())
        bulan_imunisasi.append(Imunisasi.objects.filter(tanggal__gte=start, tanggal__lte=last).count())
    
    # Detail penimbangan
    detail_penimbangan = penimbangan_list.select_related('balita').order_by('-tanggal')[:20]
    
    months = [(i, month_name[i]) for i in range(1, 13)]
    years = [2022, 2023, 2024, 2025, 2026]
    
    context = {
        'jenis_laporan': 'bulanan',
        'bulan': bulan,
        'tahun': tahun,
        'months': months,
        'years': years,
        'month_name': month_name[bulan],
        'total_balita': total_balita,
        'balita_baru': balita_baru,
        'total_penimbangan': total_penimbangan,
        'total_imunisasi': total_imunisasi,
        'balita_ditimbang': balita_ditimbang,
        'persentase_penimbangan': persentase_penimbangan,
        'status_normal': status_normal,
        'status_kurang': status_kurang,
        'status_buruk': status_buruk,
        'status_lebih': status_lebih,
        'total_status': total_status,
        'persen_normal': persen_normal,
        'persen_kurang': persen_kurang,
        'persen_buruk': persen_buruk,
        'persen_lebih': persen_lebih,
        'chart_labels': json.dumps(chart_labels),
        'chart_data': json.dumps(chart_data),
        'bulan_labels': json.dumps(bulan_labels),
        'bulan_penimbangan': json.dumps(bulan_penimbangan),
        'bulan_imunisasi': json.dumps(bulan_imunisasi),
        'detail_penimbangan': detail_penimbangan,
    }
    return render(request, 'laporan/index.html', context)


def laporan_tahunan(request, tahun):
    """Laporan tahunan"""
    start_date = date(tahun, 1, 1)
    end_date = date(tahun, 12, 31)
    
    # Total balita
    total_balita = Balita.objects.count()
    
    # Total balita baru tahun ini
    balita_baru = Balita.objects.filter(created_at__date__gte=start_date, created_at__date__lte=end_date).count()
    
    # Data per bulan
    bulan_data = []
    total_penimbangan_tahun = 0
    total_imunisasi_tahun = 0
    
    for bulan in range(1, 13):
        start = date(tahun, bulan, 1)
        last_day = monthrange(tahun, bulan)[1]
        end = date(tahun, bulan, last_day)
        
        penimbangan_bulan = Penimbangan.objects.filter(tanggal__gte=start, tanggal__lte=end).count()
        imunisasi_bulan = Imunisasi.objects.filter(tanggal__gte=start, tanggal__lte=end).count()
        
        total_penimbangan_tahun += penimbangan_bulan
        total_imunisasi_tahun += imunisasi_bulan
        
        bulan_data.append({
            'bulan': bulan,
            'nama_bulan': month_name[bulan],
            'penimbangan': penimbangan_bulan,
            'imunisasi': imunisasi_bulan,
        })
    
    # Status gizi akhir tahun
    status_normal = 0
    status_kurang = 0
    status_buruk = 0
    status_lebih = 0
    
    for balita in Balita.objects.all():
        penimbangan_terakhir = balita.penimbangan.filter(tanggal__lte=end_date).order_by('-tanggal').first()
        if penimbangan_terakhir and penimbangan_terakhir.status_gizi:
            status = penimbangan_terakhir.status_gizi
            if 'Normal' in status:
                status_normal += 1
            elif 'Kurang' in status:
                status_kurang += 1
            elif 'Buruk' in status:
                status_buruk += 1
            elif 'Lebih' in status:
                status_lebih += 1
    
    total_status = status_normal + status_kurang + status_buruk + status_lebih
    persen_normal = round((status_normal / total_status * 100), 1) if total_status > 0 else 0
    persen_kurang = round((status_kurang / total_status * 100), 1) if total_status > 0 else 0
    persen_buruk = round((status_buruk / total_status * 100), 1) if total_status > 0 else 0
    persen_lebih = round((status_lebih / total_status * 100), 1) if total_status > 0 else 0
    
    # Data untuk chart
    bulan_names = [month_name[i][:3] for i in range(1, 13)]
    penimbangan_data = [bd['penimbangan'] for bd in bulan_data]
    imunisasi_data = [bd['imunisasi'] for bd in bulan_data]
    
    # Penimbangan terbanyak
    max_penimbangan = max(penimbangan_data) if penimbangan_data else 0
    bulan_terbanyak = bulan_names[penimbangan_data.index(max_penimbangan)] if max_penimbangan > 0 else '-'
    
    context = {
        'jenis_laporan': 'tahunan',
        'tahun': tahun,
        'years': [2022, 2023, 2024, 2025, 2026],
        'total_balita': total_balita,
        'balita_baru': balita_baru,
        'total_penimbangan': total_penimbangan_tahun,
        'total_imunisasi': total_imunisasi_tahun,
        'rata_penimbangan_per_bulan': round(total_penimbangan_tahun / 12, 1),
        'rata_imunisasi_per_bulan': round(total_imunisasi_tahun / 12, 1),
        'max_penimbangan': max_penimbangan,
        'bulan_terbanyak': bulan_terbanyak,
        'status_normal': status_normal,
        'status_kurang': status_kurang,
        'status_buruk': status_buruk,
        'status_lebih': status_lebih,
        'total_status': total_status,
        'persen_normal': persen_normal,
        'persen_kurang': persen_kurang,
        'persen_buruk': persen_buruk,
        'persen_lebih': persen_lebih,
        'bulan_data': bulan_data,
        'bulan_labels': json.dumps(bulan_names),
        'penimbangan_data': json.dumps(penimbangan_data),
        'imunisasi_data': json.dumps(imunisasi_data),
    }
    return render(request, 'laporan/tahunan.html', context)


@login_required
def laporan_export_excel(request):
    """Export laporan ke Excel (Bulanan atau Tahunan)"""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        from django.http import HttpResponse
        return HttpResponse('Library openpyxl tidak terinstall. Jalankan: pip install openpyxl', status=500)
    
    from django.http import HttpResponse
    from datetime import datetime
    from calendar import monthrange, month_name
    
    jenis = request.GET.get('jenis', 'bulanan')
    bulan = int(request.GET.get('bulan', datetime.now().month))
    tahun = int(request.GET.get('tahun', datetime.now().year))
    
    from datetime import date
    
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Styles
    title_font = Font(size=16, bold=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="a93620", end_color="a93620", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    if jenis == 'tahunan':
        # ==================== LAPORAN TAHUNAN ====================
        start_date = date(tahun, 1, 1)
        end_date = date(tahun, 12, 31)
        
        total_balita = Balita.objects.count()
        total_penimbangan = Penimbangan.objects.filter(tanggal__gte=start_date, tanggal__lte=end_date).count()
        total_imunisasi = Imunisasi.objects.filter(tanggal__gte=start_date, tanggal__lte=end_date).count()
        
        ws.title = f"Laporan Tahunan {tahun}"
        
        # Title
        ws.merge_cells('A1:G1')
        ws['A1'] = f"LAPORAN TAHUNAN POSYANDU SIPOS"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws.merge_cells('A2:G2')
        ws['A2'] = f"Tahun: {tahun}"
        ws['A2'].alignment = Alignment(horizontal='center')
        
        # Summary
        ws['A4'] = "A. RINGKASAN TAHUNAN"
        ws['A4'].font = Font(bold=True)
        ws['A5'] = "Total Balita Terdaftar"
        ws['B5'] = total_balita
        ws['A6'] = "Total Penimbangan"
        ws['B6'] = total_penimbangan
        ws['A7'] = "Total Imunisasi"
        ws['B7'] = total_imunisasi
        ws['A8'] = "Rata-rata Penimbangan per Bulan"
        ws['B8'] = round(total_penimbangan / 12, 1)
        ws['A9'] = "Rata-rata Imunisasi per Bulan"
        ws['B9'] = round(total_imunisasi / 12, 1)
        
        # Data per bulan
        ws['A11'] = "B. DATA PER BULAN"
        ws['A11'].font = Font(bold=True)
        
        headers = ['Bulan', 'Penimbangan', 'Imunisasi']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=12, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        
        row = 13
        for bulan_num in range(1, 13):
            start = date(tahun, bulan_num, 1)
            last_day = monthrange(tahun, bulan_num)[1]
            end = date(tahun, bulan_num, last_day)
            
            penimbangan = Penimbangan.objects.filter(tanggal__gte=start, tanggal__lte=end).count()
            imunisasi = Imunisasi.objects.filter(tanggal__gte=start, tanggal__lte=end).count()
            
            ws.cell(row=row, column=1, value=month_name[bulan_num]).border = border
            ws.cell(row=row, column=2, value=penimbangan).border = border
            ws.cell(row=row, column=3, value=imunisasi).border = border
            row += 1
        
        # Adjust column widths
        for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
            ws.column_dimensions[col].width = 20
        
        filename = f"laporan_tahunan_{tahun}.xlsx"
        
    else:
        # ==================== LAPORAN BULANAN ====================
        start_date = date(tahun, bulan, 1)
        last_day = monthrange(tahun, bulan)[1]
        end_date = date(tahun, bulan, last_day)
        
        total_balita = Balita.objects.count()
        penimbangan_list = Penimbangan.objects.filter(tanggal__gte=start_date, tanggal__lte=end_date)
        total_penimbangan = penimbangan_list.count()
        
        ws.title = f"Laporan {month_name[bulan]} {tahun}"
        
        # Title
        ws.merge_cells('A1:F1')
        ws['A1'] = f"LAPORAN BULANAN POSYANDU SIPOS"
        ws['A1'].font = title_font
        ws['A1'].alignment = Alignment(horizontal='center')
        
        ws.merge_cells('A2:F2')
        ws['A2'] = f"Bulan: {month_name[bulan]} {tahun}"
        ws['A2'].alignment = Alignment(horizontal='center')
        
        # Summary
        ws['A4'] = "A. RINGKASAN"
        ws['A4'].font = Font(bold=True)
        ws['A5'] = "Total Balita Terdaftar"
        ws['B5'] = total_balita
        ws['A6'] = "Total Penimbangan Bulan Ini"
        ws['B6'] = total_penimbangan
        ws['A7'] = "Total Imunisasi Bulan Ini"
        ws['B7'] = Imunisasi.objects.filter(tanggal__gte=start_date, tanggal__lte=end_date).count()
        
        # Status Gizi
        ws['A9'] = "B. STATUS GIZI"
        ws['A9'].font = Font(bold=True)
        
        headers = ['Status Gizi', 'Jumlah', 'Persentase']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=10, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        
        status_data = []
        for balita in Balita.objects.all():
            penimbangan_terakhir = balita.penimbangan.filter(tanggal__lte=end_date).order_by('-tanggal').first()
            if penimbangan_terakhir and penimbangan_terakhir.status_gizi:
                status = penimbangan_terakhir.status_gizi
                if 'Normal' in status:
                    status_data.append('Normal')
                elif 'Kurang' in status:
                    status_data.append('Gizi Kurang')
                elif 'Buruk' in status:
                    status_data.append('Gizi Buruk')
                elif 'Lebih' in status:
                    status_data.append('Gizi Lebih')
        
        from collections import Counter
        status_count = Counter(status_data)
        total = len(status_data)
        
        row = 11
        for status, count in status_count.items():
            percentage = round(count / total * 100, 1) if total > 0 else 0
            ws.cell(row=row, column=1, value=status).border = border
            ws.cell(row=row, column=2, value=count).border = border
            ws.cell(row=row, column=3, value=f"{percentage}%").border = border
            row += 1
        
        # Detail Penimbangan
        ws.cell(row=row + 2, column=1, value="C. DETAIL PENIMBANGAN").font = Font(bold=True)
        
        detail_headers = ['Tanggal', 'Nama Balita', 'Berat Badan (kg)', 'Tinggi Badan (cm)', 'Status Gizi']
        for col, header in enumerate(detail_headers, 1):
            cell = ws.cell(row=row + 3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = border
        
        r = row + 4
        for p in penimbangan_list[:50]:
            ws.cell(row=r, column=1, value=p.tanggal.strftime('%d/%m/%Y')).border = border
            ws.cell(row=r, column=2, value=p.balita.nama).border = border
            ws.cell(row=r, column=3, value=p.berat_badan).border = border
            ws.cell(row=r, column=4, value=p.tinggi_badan).border = border
            ws.cell(row=r, column=5, value=p.status_gizi if p.status_gizi else '-').border = border
            r += 1
        
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            ws.column_dimensions[col].width = 20
        
        filename = f"laporan_{month_name[bulan]}_{tahun}.xlsx"
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def laporan_export_pdf(request):
    """Export laporan ke PDF (Bulanan atau Tahunan)"""
    try:
        from django.template.loader import get_template
        from xhtml2pdf import pisa
        import io
    except ImportError:
        from django.http import HttpResponse
        return HttpResponse('Library xhtml2pdf tidak terinstall. Jalankan: pip install xhtml2pdf', status=500)
    
    from django.http import HttpResponse
    from datetime import datetime
    from calendar import monthrange, month_name
    
    jenis = request.GET.get('jenis', 'bulanan')
    bulan = int(request.GET.get('bulan', datetime.now().month))
    tahun = int(request.GET.get('tahun', datetime.now().year))
    
    from datetime import date
    start_date = date(tahun, bulan, 1)
    last_day = monthrange(tahun, bulan)[1]
    end_date = date(tahun, bulan, last_day)
    
    # Data umum
    total_balita = Balita.objects.count()
    
    # Status gizi
    status_normal = 0
    status_kurang = 0
    status_buruk = 0
    status_lebih = 0
    
    for balita in Balita.objects.all():
        penimbangan_terakhir = balita.penimbangan.filter(tanggal__lte=end_date).order_by('-tanggal').first()
        if penimbangan_terakhir and penimbangan_terakhir.status_gizi:
            status = penimbangan_terakhir.status_gizi
            if 'Normal' in status:
                status_normal += 1
            elif 'Kurang' in status:
                status_kurang += 1
            elif 'Buruk' in status:
                status_buruk += 1
            elif 'Lebih' in status:
                status_lebih += 1
    
    total_status = status_normal + status_kurang + status_buruk + status_lebih
    persen_normal = round((status_normal / total_status * 100), 1) if total_status > 0 else 0
    persen_kurang = round((status_kurang / total_status * 100), 1) if total_status > 0 else 0
    persen_buruk = round((status_buruk / total_status * 100), 1) if total_status > 0 else 0
    persen_lebih = round((status_lebih / total_status * 100), 1) if total_status > 0 else 0
    
    if jenis == 'tahunan':
        # Laporan Tahunan
        penimbangan_list = Penimbangan.objects.filter(tanggal__gte=date(tahun, 1, 1), tanggal__lte=date(tahun, 12, 31))
        total_penimbangan = penimbangan_list.count()
        total_imunisasi = Imunisasi.objects.filter(tanggal__gte=date(tahun, 1, 1), tanggal__lte=date(tahun, 12, 31)).count()
        
        context = {
            'jenis': 'tahunan',
            'tahun': tahun,
            'total_balita': total_balita,
            'total_penimbangan': total_penimbangan,
            'total_imunisasi': total_imunisasi,
            'status_normal': status_normal,
            'status_kurang': status_kurang,
            'status_buruk': status_buruk,
            'status_lebih': status_lebih,
            'persen_normal': persen_normal,
            'persen_kurang': persen_kurang,
            'persen_buruk': persen_buruk,
            'persen_lebih': persen_lebih,
            'total_status': total_status,
            'now': datetime.now(),
        }
        template = get_template('laporan/pdf_tahunan.html')
    else:
        # Laporan Bulanan
        penimbangan_list = Penimbangan.objects.filter(tanggal__gte=start_date, tanggal__lte=end_date)
        total_penimbangan = penimbangan_list.count()
        total_imunisasi = Imunisasi.objects.filter(tanggal__gte=start_date, tanggal__lte=end_date).count()
        
        context = {
            'jenis': 'bulanan',
            'bulan': bulan,
            'tahun': tahun,
            'month_name': month_name[bulan],
            'total_balita': total_balita,
            'total_penimbangan': total_penimbangan,
            'total_imunisasi': total_imunisasi,
            'status_normal': status_normal,
            'status_kurang': status_kurang,
            'status_buruk': status_buruk,
            'status_lebih': status_lebih,
            'persen_normal': persen_normal,
            'persen_kurang': persen_kurang,
            'persen_buruk': persen_buruk,
            'persen_lebih': persen_lebih,
            'detail_penimbangan': penimbangan_list[:30],
            'now': datetime.now(),
        }
        template = get_template('laporan/pdf.html')
    
    html = template.render(context)
    result = io.BytesIO()
    pdf = pisa.pisaDocument(io.BytesIO(html.encode("UTF-8")), result)
    
    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        if jenis == 'tahunan':
            response['Content-Disposition'] = f'attachment; filename="laporan_tahunan_{tahun}.pdf"'
        else:
            response['Content-Disposition'] = f'attachment; filename="laporan_{month_name[bulan]}_{tahun}.pdf"'
        return response
    
    return HttpResponse('Error generating PDF', status=400)